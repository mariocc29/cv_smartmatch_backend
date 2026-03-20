from django.http import HttpResponse
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt

from web.services.document.components.document_generator import DocumentGenerator


class DOCGenerator(DocumentGenerator):
  def __init__(self, response: HttpResponse):
    self.response = response

  def content_type(self) -> str:
    return 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

  def generate_document(self, rendered_html: str) -> HttpResponse:
    soup = BeautifulSoup(rendered_html, 'html.parser')
    document = Document()
    
    for div in soup.find_all(recursive=True):
      text = div.text.strip()
      div_class = div.get('class', [])

      p = document.add_paragraph('')
      p.alignment = 3
      p.paragraph_format.font.size = Pt(13)
      p.paragraph_format.space_after = 0
      run = p.add_run(text)
      
      if 'bold' in div_class:
        run.bold = True
      
      if 'align-center' in div_class:
        p.alignment = 1

      if 'text-large' in div_class:
        p.font.size = Pt(14)
      
      if 'text-small' in div_class:
        p.font.size = Pt(12)
      
      if 'text-small' in div_class:
        p.font.caption = Pt(10)

    document.save(self.response)